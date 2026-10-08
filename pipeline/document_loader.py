from pathlib import Path
from docx import Document


SUPPORTED_TEXT = {".docx"}
SUPPORTED_IMAGE = {".png", ".jpg", ".jpeg"}


def read_docx(path: Path) -> str:
    doc = Document(str(path))
    chunks = []

    for paragraph in doc.paragraphs:
        value = paragraph.text.strip()
        if value:
            chunks.append(value)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            row_text = " | ".join(x for x in cells if x)
            if row_text:
                chunks.append(row_text)

    return "\n".join(chunks)


def load_document(path: str | Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix in SUPPORTED_TEXT:
        return read_docx(path)

    if suffix in SUPPORTED_IMAGE:
        from pipeline.image_reader import ocr_image
        return ocr_image(path)

    raise ValueError(f"Unsupported document type: {path.suffix}")
